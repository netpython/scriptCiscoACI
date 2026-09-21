#!/usr/bin/env python3
"""Audit structurel des tenants : VRF, BD, EPG, subnets et contrats."""
from __future__ import annotations
import re, sys
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'scripts'))
from common import *
log = setup('tenant_policy_compliance')
a = parser('Conformité tenants ACI').parse_args()
cfg = config(a.config)
objects = {}
with APIC(cfg) as apic:
    for cls in ('fvTenant', 'fvCtx', 'fvBD', 'fvAEPg', 'fvSubnet', 'vzBrCP', 'fvRsCtx', 'fvRsBd'):
        objects[cls] = rows(apic.cls(cls))
tenants = []
for t in objects['fvTenant']:
    name = t.get('name')
    prefix = f'uni/tn-{name}/'
    counts = {k: sum((str(x.get('dn', '')).startswith(prefix) for x in v)) for k, v in objects.items() if k != 'fvTenant'}
    checks = {'has_vrf': counts['fvCtx'] > 0, 'has_bd': counts['fvBD'] > 0, 'has_epg': counts['fvAEPg'] > 0, 'bd_has_vrf': counts['fvRsCtx'] >= counts['fvBD'], 'epg_has_bd': counts['fvRsBd'] >= counts['fvAEPg']}
    score = round(100 * sum(checks.values()) / len(checks))
    tenants.append({'tenant': name, **counts, **{f'check_{k}': 'PASS' if v else 'FAIL' for k, v in checks.items()}, 'score': score, 'compliance': 'OK' if score == 100 else 'WARNING' if score >= 60 else 'CRITICAL'})
wb = Workbook()
ws = wb.active
ws.title = 'Tenant Compliance'
headers = list(tenants[0]) if tenants else ['result']
ws.append(headers)
for r in tenants:
    ws.append([r.get(h, '') for h in headers])
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor='1F4E78')
ws.freeze_panes = 'A2'
path = OUTPUT / f'tenant_policy_compliance_{stamp()}.xlsx'
wb.save(path)
log.info('%s', path)
